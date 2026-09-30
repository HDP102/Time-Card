"""
sc_recert_submit.py - create RLM recertification tickets through the SecureChange REST
API and walk them to Closed, with each rule's owner and decision recorded in the ticket.

Why this exists
    recert_processor.py uses SecureTrack GraphQL `createTicketDraft`, which only makes
    Drafts; drafts cannot be submitted by API (the Requests API has search + cancel only).
    This script creates real tickets instead, one per ledger batch, so ticket N covers
    the same rules as draft N. Cancel the drafts once the tickets are confirmed.

How a ticket moves (RLM Rule Recertification Workflow, from ticket 139744)
    1 Open Recertification Request   devices -> bindings -> rules          (POST)
    2 Recertification Decision       + rule_recertification_info per rule  (PUT task, DONE)
    3 Update Recertification Data    update_metadata, then DONE            -> Closed

Identifiers SecureChange needs come from SecureTrack REST, by device name + rule name:
    device id (management_id), rule uid {GUID}, binding uid. Cached per device.

Owner per rule: first CorpID found in OWNER_ROLES order in the reporting views
("Last, First (CORPID)"), email CORPID@EMAIL_DOMAIN.

Run from RecertProcessor/:   py sc_recert_submit.py
Start with DRY_RUN = True, then one small live ticket, then the batches Swapnil asked for.
"""
import getpass
import json
import os
import re
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import urlparse

import requests
import urllib3
from openpyxl import load_workbook

import recert_processor as r

urllib3.disable_warnings()

# ═══════════════════════════ CONFIG ═══════════════════════════
SOURCE = "ledger"         # where the rules come from:
                          #   "ledger"  - recert_processor.py batches (submitted_rules.jsonl), 1-123
                          #   "reports" - Tufin rule reports in SharedData (RULE_REPORT_GLOB):
                          #               every rule in them is recertified, no decision column.
                          #               Batches are numbered from 1001 so they never collide.
                          #   "sweep"   - BOTH of the above: every rule from the ledger and the rule
                          #               reports that is not in a ticket yet, packed into tickets of
                          #               up to 300. Numbered from 2001 (continues after earlier sweeps).
                          #               Use it to pick up leftovers without one tiny ticket per batch.
DRY_RUN = True            # True: resolve everything, save the payload, change nothing
STOP_AFTER = "create"     # "create"   - make the ticket, stop (check it in SecureChange)
                          # "decision" - also fill + complete step 2
                          # "close"    - walk it all the way to Closed
BATCHES = [1]             # batch numbers; for "reports" 77 and 1077 both mean batch 1077. Examples:
                          #   [1, 2, 3, 4, 5]          specific batches
                          #   list(range(6, 124))      a range: 6 through 123 (end is exclusive)
                          #   "all"                    every batch in the ledger
                          # Batches already fully ticketed are skipped, so re-running a range is safe.
TEST_RULE_LIMIT = 5       # cap rules per ticket for the first test; 0 = full batch

WORKFLOW_NAME = "RLM Rule Recertification Workflow"
SUBJECT = "Recertify Firewall Rules - NPS - batch {batch}"
PRIORITY = "Normal"
EXPIRY_DAYS = 365                     # Megan's sample used ~1 year
EMAIL_DOMAIN = "pge.com"
DESCRIPTION = "Recertified per owner decision in {sheet} ({date}). NPS recertification automation."
OWNER_ROLES = ["Client Owner", "IT SME", "IT Lead", "IT SME Backup", "IT Lead Delegate"]
FALLBACK_OWNER = ""       # CorpID to use when a rule has no owner; "" = leave owner blank

# "reports" source (Tufin rule report exports, header row found automatically)
RULE_REPORT_GLOB = "*Rule_report*.xlsx"
REPORT_BATCH_BASE = 1000
SWEEP_BATCH_BASE = 2000
REPORT_OWNER_COLUMNS = ["Business Owner", "Application Owner"]   # first CorpID found wins
DESCRIPTION_REPORTS = "Recertified per NPS recertification list {sheet} ({date}). NPS recertification automation."

STEP_OPEN = "Open Recertification Request"
STEP_DECISION = "Recertification Decision"
STEP_UPDATE = "Update Recertification Data"
# ══════════════════════════════════════════════════════════════

HOST = urlparse(r.SECURETRACK_GRAPHQL).netloc
SC = f"https://{HOST}/securechangeworkflow/api/securechange"
ST = f"https://{HOST}/securetrack/api"
OUT = Path(r.OUTPUT_DIR)
CACHE_DIR = OUT / "st_rules_cache"
SC_LEDGER = OUT / "sc_tickets.jsonl"
JSON_HDRS = {"Accept": "application/json", "Content-Type": "application/json"}
CORPID_RE = re.compile(r"\(([A-Za-z0-9]{3,8})\)\s*$")

AUTH = None


def L(x):
    return [] if x is None else (x if isinstance(x, list) else [x])


def log(msg):
    print(f"{datetime.now():%H:%M:%S} {msg}", flush=True)


def call(method, url, **kw):
    """HTTP with Tufin's error text surfaced. Returns the response; raises on non-2xx."""
    resp = requests.request(method, url, auth=AUTH, headers=JSON_HDRS, verify=False,
                            timeout=300, **kw)
    if not 200 <= resp.status_code < 300:
        text = resp.text
        m = re.search(r"<message>(.*?)</message>", text, re.S)
        raise RuntimeError(f"{method} {url} -> {resp.status_code}: {m.group(1) if m else text[:800]}")
    return resp


# ─────────────────────────── inputs ───────────────────────────

def ledger_batches():
    """Recertify batches from recert_processor's ledger, in order. Batch N = draft N."""
    out = []
    for line in open(r.LEDGER_PATH, encoding="utf-8"):
        if line.strip():
            rec = json.loads(line)
            if rec.get("workflow_type") == r.RECERTIFY_WORKFLOW_TYPE:
                out.append(rec["uids"])
    return out


def open_wb(path):
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        raise SystemExit(f"{Path(path).name} is open in Excel (or OneDrive is syncing it) - "
                         f"close it and run again.")


def mapping_by_uid():
    """GraphQL uid -> (device name, rule name), from the Rule-ID mapping workbook."""
    wb = open_wb(r.data_path(r.MAPPING_PATH))
    ws = next((wb[n] for n in wb.sheetnames
               if n.strip().lower() == r.MAPPING_SHEET.lower()), wb.active)
    rows = ws.iter_rows(values_only=True)
    h = [str(x).strip() if x is not None else "" for x in next(rows)]
    d, n, u = h.index(r.MAP_DEVICE_COLUMN), h.index(r.MAP_RULE_COLUMN), h.index(r.MAP_UID_COLUMN)
    out = {}
    for row in rows:
        if row[u]:
            out[str(row[u]).strip()] = (str(row[d] or "").strip(), str(row[n] or "").strip())
    wb.close()
    return out


def owners_by_rule():
    """(DEVICE||RULE) -> (corpid, role, sheet) from the reporting views."""
    out = {}
    for path in r.resolve_inputs(r.REPORT_PATH):
        wb = open_wb(path)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [str(x).strip() if x is not None else "" for x in next(rows, [])]
        di, _ = r.find_column(header, r.DEVICE_COLUMN)
        ri, _ = r.find_column(header, r.RULE_NAME_COLUMN)
        roles = [(role, header.index(role)) for role in OWNER_ROLES if role in header]
        if di is None or ri is None or not roles:
            log(f"  {path.name}: no device/rule/owner columns - skipped")
            wb.close()
            continue
        for row in rows:
            k = r._key(row[di], row[ri])
            if k in out:
                continue
            for role, i in roles:
                m = CORPID_RE.search(str(row[i] or ""))
                if m:
                    out[k] = (m.group(1).upper(), role, path.name)
                    break
        wb.close()
    return out


_CORPID_TOKEN = re.compile(r"\b[A-Za-z0-9]{4}\b")


def report_batches():
    """Rules from the Tufin rule reports -> (batches, mapping, owners), same shapes as the
    ledger source. Every rule is recertified. A rule in more than one file is kept once."""
    files = sorted(r.DATA_DIR.glob(RULE_REPORT_GLOB))
    if not files:
        raise SystemExit(f"No files matching {RULE_REPORT_GLOB!r} in {r.DATA_DIR}")
    order, mapping, owners = [], {}, {}
    for f in files:
        wb = open_wb(f)
        rows = wb[wb.sheetnames[0]].iter_rows(values_only=True)
        header = None
        for _ in range(20):                      # CSV exports have metadata lines on top
            row = next(rows, None)
            if row is None:
                break
            cells = [str(c).strip() if c is not None else "" for c in row]
            if "SecureTrack Rule ID" in cells and "Rule Name" in cells:
                header = cells
                break
        if not header:
            log(f"  {f.name}: no 'SecureTrack Rule ID' / 'Rule Name' header - skipped")
            wb.close()
            continue
        ui, ni = header.index("SecureTrack Rule ID"), header.index("Rule Name")
        di = header.index("Device Name") if "Device Name" in header else None
        ocols = [(c, header.index(c)) for c in REPORT_OWNER_COLUMNS if c in header]
        added = dup = 0
        for row in rows:
            uid = str(row[ui] or "").strip() if ui < len(row) else ""
            if not uid:
                continue
            if uid in mapping:
                dup += 1
                continue
            dev = str(row[di] or "").strip() if di is not None else ""
            name = str(row[ni] or "").strip()
            mapping[uid] = (dev, name)
            order.append(uid)
            added += 1
            for col, i in ocols:
                m = _CORPID_TOKEN.search(str(row[i] or ""))
                if m:
                    owners[r._key(dev, name)] = (m.group(0).upper(), col, f.name)
                    break
        wb.close()
        log(f"  {f.name}: {added:,} rule(s){f', {dup} already in another file' if dup else ''}")
    size = r.TICKET_BATCH_SIZE
    return [order[i:i + size] for i in range(0, len(order), size)], mapping, owners


# ─────────────────────── SecureTrack lookups ───────────────────────

_devs = {}
_st_dev = {}
_UNRESOLVED = []          # details for rules that could not be resolved -> unresolved_debug.json


def st_devices(name):
    """All SecureTrack devices whose name matches exactly -> [(id, parent_id)].
    Usually one. Can be more: the same device-group name under two Panoramas
    (e.g. BishopRanch 8484 / 6894, Metcalf 6889 / 4650, Fresno Regional Office 9611 / 9543)."""
    if name not in _devs:
        data = call("GET", f"{ST}/devices", params={"name": name}).json()
        found = []
        for d in L((data.get("devices") or {}).get("device")):
            _st_dev[d["id"]] = d
            if str(d.get("name", "")).strip().lower() == name.strip().lower():
                found.append((d["id"], d.get("parent_id")))
        _devs[name] = found
        if len(found) > 1:
            log(f"  device {name!r}: {len(found)} devices share this name "
                f"({', '.join(str(i) for i, _ in found)}) - resolving per rule")
    return _devs[name]


def st_device(dev_id):
    if dev_id not in _st_dev:
        d = call("GET", f"{ST}/devices/{dev_id}").json()
        _st_dev[dev_id] = d.get("device") or d
    return _st_dev[dev_id]


def st_ancestors(dev_id):
    """Names of the device's parents in SecureTrack, nearest first (parent group ... Panorama)."""
    names, seen, cur = [], set(), st_device(dev_id).get("parent_id")
    while cur and cur not in seen and len(names) < 6:
        seen.add(cur)
        d = st_device(cur)
        names.append(str(d.get("name", "")).strip())
        cur = d.get("parent_id")
    return names


# GraphQL: the rule's device ancestry and its position in the policy. Deep parent chain first;
# falls back to one level if this Tufin version won't nest parent.
_GQL_DEEP = """query($f:String){ rules(filter:$f){ values(first:200){ id policyIndex priority
  device{ name parent{ name parent{ name parent{ name } } } } } } }"""
_GQL_SHALLOW = """query($f:String){ rules(filter:$f){ values(first:200){ id policyIndex priority
  device{ name parent{ name } } } } }"""
_gql_query = [_GQL_DEEP]
_gql_info = {}


def _chain(dev):
    names, p = [], (dev or {}).get("parent")
    while p:
        names.append(str(p.get("name") or "").strip())
        p = p.get("parent")
    return names


def graphql_rule_info(graphql_id, dev, rule_name):
    """-> {"ancestors": [...nearest first], "rank": i, "count": n} or None.
    rank/count: this rule's position among same-name rules on the same device (policy order)."""
    if graphql_id in _gql_info:
        return _gql_info[graphql_id]
    info = None
    if "'" not in rule_name:
        vals = None
        while vals is None and _gql_query:
            try:
                vals = r._graphql(_gql_query[0], {"f": f"name = '{rule_name}'"})["data"]["rules"]["values"] or []
            except Exception as e:
                if _gql_query[0] is _GQL_DEEP:
                    log(f"    GraphQL deep parent query not supported ({e}) - using one level")
                    _gql_query[0] = _GQL_SHALLOW
                else:
                    log(f"    GraphQL lookup failed for {rule_name!r}: {e}")
                    vals = []
        me = next((v for v in vals if v.get("id") == graphql_id), None)
        if me:
            anc = _chain(me.get("device"))
            dname = str((me.get("device") or {}).get("name") or "").strip().lower()
            same = [v for v in vals
                    if str((v.get("device") or {}).get("name") or "").strip().lower() == dname
                    and _chain(v.get("device")) == anc]
            order = lambda v: (v.get("policyIndex") is None, v.get("policyIndex") or 0, v.get("priority") or 0)
            same.sort(key=order)
            info = {"ancestors": anc, "count": len(same),
                    "rank": next(i for i, v in enumerate(same) if v.get("id") == graphql_id)}
    _gql_info[graphql_id] = info
    return info


def pick_device(graphql_id, dev, rule):
    """-> (dev_id, None) or (None, reason). Same-name devices are resolved by which one has
    the rule, then by the closest Panorama/parent group that tells them apart."""
    cands = st_devices(dev)
    if not cands:
        return None, f"device {dev!r} not found in SecureTrack"
    if len(cands) == 1:
        return cands[0][0], None
    has_rule = [i for i, _ in cands if st_rules(i).get(r._norm(rule))]
    if len(has_rule) == 1:
        return has_rule[0], None
    if not has_rule:
        return None, f"{dev} / {rule}: not on any of the {len(cands)} devices named {dev!r}"
    info = graphql_rule_info(graphql_id, dev, rule)
    st_anc = {i: [a.lower() for a in st_ancestors(i)] for i in has_rule}
    if info:
        for a in info["ancestors"]:
            match = [i for i in has_rule if a.lower() in st_anc[i]]
            if len(match) == 1:
                return match[0], None
    _UNRESOLVED.append({"graphql_id": graphql_id, "device": dev, "rule": rule,
                        "problem": "same-name devices",
                        "graphql_ancestors": info["ancestors"] if info else None,
                        "securetrack_candidates": {str(i): st_ancestors(i) for i in has_rule}})
    return None, (f"{dev} / {rule}: ambiguous - on {len(has_rule)} devices named {dev!r} "
                  f"({', '.join(str(i) for i in has_rule)}), Panorama not resolved")


def _binding_uid(rule):
    b = rule.get("binding")
    for item in L(b):
        if isinstance(item, dict):
            if item.get("uid"):
                return item["uid"]
            inner = item.get("binding")
            if isinstance(inner, dict) and inner.get("uid"):
                return inner["uid"]
    return rule.get("binding_uid")


_revisions = {}


def st_revision(dev_id):
    """Current policy revision id for a device - SecureChange requires it per device
    (ticket 139744 carries revision_id 33354632, the same value as each rule's version_id)."""
    if dev_id not in _revisions:
        data = call("GET", f"{ST}/devices/{dev_id}/latest_revision").json()
        rev = data.get("revision", data)
        rev = rev[0] if isinstance(rev, list) else rev
        _revisions[dev_id] = rev.get("id")
        log(f"  device {dev_id}: latest revision id {_revisions[dev_id]} "
            f"(revision number {rev.get('revisionId', '?')})")
    return _revisions[dev_id]


_rule_idx = {}


def st_rules(dev_id):
    """rule name (normalised) -> [[uid, binding_uid, order], ...] for one device, in policy
    order. Cached to disk."""
    if dev_id in _rule_idx:
        return _rule_idx[dev_id]
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cf = CACHE_DIR / f"{dev_id}.v2.json"       # v2: [uid, binding, rule order]
    if cf.exists() and time.time() - cf.stat().st_mtime < 24 * 3600:
        idx = json.loads(cf.read_text(encoding="utf-8"))
    else:
        log(f"  fetching rules for device {dev_id} from SecureTrack (can take a minute)...")
        data = call("GET", f"{ST}/devices/{dev_id}/rules").json()
        idx = {}
        for x in L((data.get("rules") or {}).get("rule")):
            order = x.get("rule_number", x.get("order", x.get("number")))
            idx.setdefault(r._norm(x.get("name")), []).append([x.get("uid"), _binding_uid(x), order])
        cf.write_text(json.dumps(idx), encoding="utf-8")
    _rule_idx[dev_id] = idx
    return idx


# ─────────────────────────── payloads ───────────────────────────

def resolve_batch(uids, mapping, owners):
    """-> (resolved rules, problems). Each rule: dev_id, uid, binding, owner, sheet."""
    ok, problems = [], []
    for g in uids:
        if g not in mapping:
            problems.append((g, "not in mapping workbook"))
            continue
        dev, rule = mapping[g]
        dev_id, why = pick_device(g, dev, rule)
        if dev_id is None:
            problems.append((g, why))
            continue
        hits = st_rules(dev_id).get(r._norm(rule), [])
        how = "name"
        if len(hits) > 1:
            # Same rule name more than once on the device (e.g. pre- and post-rulebase).
            # Pick by position: GraphQL gives this rule's rank among the same-name rules on
            # the device, SecureTrack lists them in policy order. Only if both counts agree.
            info = graphql_rule_info(g, dev, rule)
            if all(isinstance(h[2], (int, float)) for h in hits):
                hits = sorted(hits, key=lambda h: h[2])
            if info and info["count"] == len(hits):
                hits, how = [hits[info["rank"]]], f"position {info['rank'] + 1} of {info['count']}"
            else:
                _UNRESOLVED.append({"graphql_id": g, "device": dev, "rule": rule, "device_id": dev_id,
                                    "problem": "duplicate rule name on device",
                                    "securetrack_matches": len(hits),
                                    "graphql_same_name_on_device": info["count"] if info else None})
        if len(hits) != 1:
            problems.append((g, f"{dev} / {rule}: {len(hits)} matches on device {dev_id}"))
            continue
        uid, binding = hits[0][0], hits[0][1]
        if not uid or not binding:
            problems.append((g, f"{dev} / {rule}: missing uid or binding"))
            continue
        rev_id = st_revision(dev_id)
        if not rev_id:
            problems.append((g, f"device {dev_id}: no latest revision in SecureTrack"))
            continue
        owner = owners.get(r._key(dev, rule))
        ok.append({"graphql_id": g, "device": dev, "rule": rule, "dev_id": int(dev_id),
                   "revision_id": int(rev_id),
                   "uid": uid, "binding": binding, "resolved_by": how,
                   "owner": owner[0] if owner else FALLBACK_OWNER,
                   "owner_role": owner[1] if owner else "",
                   "sheet": owner[2] if owner else ""})
    return ok, problems


def devices_block(rules):
    by_dev, revs = {}, {}
    for x in rules:
        by_dev.setdefault(x["dev_id"], {}).setdefault(x["binding"], []).append(x["uid"])
        revs[x["dev_id"]] = x["revision_id"]
    return {"device": [
        {"management_id": dev_id,
         "revision_id": revs[dev_id],
         "bindings": {"binding": [
             {"binding_uid": b, "rules": {"rule": [{"uid": u} for u in uids]}}
             for b, uids in bindings.items()]}}
        for dev_id, bindings in by_dev.items()]}


def infos_block(rules):
    today = date.today().isoformat()
    expiry = (date.today() + timedelta(days=EXPIRY_DAYS)).isoformat()
    infos = []
    for x in rules:
        info = {"@xsi.type": "certify", "binding_uid": x["binding"], "rule_uid": x["uid"],
                "description": (DESCRIPTION if SOURCE == "ledger" or
                                (SOURCE == "sweep" and "Rule_report" not in (x["sheet"] or ""))
                                else DESCRIPTION_REPORTS).format(sheet=x["sheet"] or "reporting view", date=today),
                "certification_expiration_date": expiry}
        if x["owner"]:
            info["business_owner"] = x["owner"]
            info["business_owner_email"] = f"{x['owner']}@{EMAIL_DOMAIN}"
        infos.append(info)
    return {"rule_recertification_info": infos}


def create_payload(batch_no, rules):
    return {"ticket": {
        "subject": SUBJECT.format(batch=batch_no),
        "priority": PRIORITY,
        "workflow": {"name": WORKFLOW_NAME},
        "steps": {"step": [{"name": STEP_OPEN, "tasks": {"task": [{"fields": {"field": [{
            "@xsi.type": "rule_recertification",
            "name": "Recertification Field",
            "devices": devices_block(rules),
        }]}}]}}]},
    }}


# ─────────────────────────── SecureChange ───────────────────────────

def sc_ledger(rec):
    rec["ts"] = datetime.now().isoformat(timespec="seconds")
    with open(SC_LEDGER, "a", encoding="utf-8") as f:
        f.write(json.dumps(rec) + "\n")
        f.flush()
        os.fsync(f.fileno())


def sc_history():
    """Read the SC ledger. Returns (ticketed, open_tickets):
        ticketed      - every GraphQL rule id already in ANY ticket we created
        open_tickets  - {ticket_id: {"batch", "uids", "stage"}} for tickets not yet closed
    Tracking by rule (not by batch) means a small test ticket and the rest of the
    batch never overlap, and no rule is ever put in a second ticket."""
    tickets = {}
    if SC_LEDGER.exists():
        for line in open(SC_LEDGER, encoding="utf-8"):
            if line.strip():
                rec = json.loads(line)
                t = tickets.setdefault(rec["ticket_id"], {"batch": rec["batch"], "uids": []})
                if rec.get("uids"):
                    t["uids"] = rec["uids"]
                t["stage"] = rec["stage"]
    ticketed = {u for t in tickets.values() for u in t["uids"]}
    open_tickets = {tid: t for tid, t in tickets.items() if t["stage"] != "closed"}
    return ticketed, open_tickets


def sc_batch_numbers():
    """Every batch number recorded in the ticket ledger (so sweep numbers never repeat)."""
    nums = set()
    if SC_LEDGER.exists():
        for line in open(SC_LEDGER, encoding="utf-8"):
            if line.strip():
                nums.add(json.loads(line)["batch"])
    return nums


def get_ticket(tid):
    return call("GET", f"{SC}/tickets/{tid}").json()["ticket"]


def current(t):
    """(step name, task id, field dict) for the ticket's current step."""
    name = (t.get("current_step") or {}).get("name")
    for step in L((t.get("steps") or {}).get("step")):
        if step.get("name") == name:
            task = L((step.get("tasks") or {}).get("task"))[0]
            field = L((task.get("fields") or {}).get("field"))[0]
            return name, task["id"], field
    return name, None, None


def complete_task(tid, task_id, field):
    body = {"task": {"id": task_id, "status": "DONE", "fields": {"field": [field]}}}
    call("PUT", f"{SC}/tickets/{tid}/steps/current/tasks/{task_id}", data=json.dumps(body))


def advance(tid, batch_no, rules):
    """Walk the ticket forward until STOP_AFTER or Closed."""
    for _ in range(8):
        t = get_ticket(tid)
        status = t.get("status", "")
        step, task_id, field = current(t)
        log(f"  ticket {tid}: status={status!r} step={step!r}")
        if "closed" in status.lower() or step is None:
            sc_ledger({"batch": batch_no, "ticket_id": tid, "stage": "closed"})
            return True
        if step == STEP_OPEN:
            complete_task(tid, task_id, field)
        elif step == STEP_DECISION:
            field["rule_recertification_infos"] = infos_block(rules)
            complete_task(tid, task_id, field)
            sc_ledger({"batch": batch_no, "ticket_id": tid, "stage": "decision"})
            if STOP_AFTER == "decision":
                log("  STOP_AFTER='decision' - stopping here")
                return True
        elif step == STEP_UPDATE:
            call("PUT", f"{SC}/tickets/{tid}/steps/current/tasks/{task_id}"
                        f"/rule_recertification/update_metadata")
            time.sleep(3)
            _, task_id, field = current(get_ticket(tid))
            if task_id:
                complete_task(tid, task_id, field)
        else:
            log(f"  unexpected step {step!r} - stopping so a person can look")
            return False
        time.sleep(2)
    log("  ticket did not close within 8 steps - check it in SecureChange")
    return False


# ─────────────────────────── main ───────────────────────────

def main():
    global AUTH
    user = os.environ.get("SECURETRACK_USER") or input("Tufin username: ").strip()
    pwd = os.environ.get("SECURETRACK_PASSWORD") or getpass.getpass("Tufin password: ")
    AUTH = (user, pwd)
    # Same login for the GraphQL lookup used to tell same-name devices apart (no second prompt).
    r._TOKEN_CACHE["username"], r._TOKEN_CACHE["password"] = user, pwd

    log(f"Host {HOST} | DRY_RUN={DRY_RUN} | STOP_AFTER={STOP_AFTER} | "
        f"BATCHES={BATCHES if BATCHES == 'all' or len(BATCHES) < 8 else f'{BATCHES[0]}..{BATCHES[-1]} ({len(BATCHES)})'} | TEST_RULE_LIMIT={TEST_RULE_LIMIT}")
    if SOURCE == "reports":
        log(f"Source: rule reports ({RULE_REPORT_GLOB}) - every rule is recertified")
        batches, mapping, owners = report_batches()
        base = REPORT_BATCH_BASE
        log(f"  {len(mapping):,} rules in {len(batches)} batch(es) "
            f"(numbered {base + 1}-{base + len(batches)}), {len(owners):,} with an owner")
    elif SOURCE == "sweep":
        log("Source: sweep - every ledger and rule-report rule not in a ticket yet")
        lb = ledger_batches()
        mapping, owners = mapping_by_uid(), owners_by_rule()
        log(f"  ledger (submitted_rules.jsonl): {sum(len(b) for b in lb):,} rule(s) in {len(lb)} batch(es)")
        rb = []
        if list(r.DATA_DIR.glob(RULE_REPORT_GLOB)):
            rb, rmap, rown = report_batches()
            for k, v in rmap.items():
                mapping.setdefault(k, v)
            for k, v in rown.items():
                owners.setdefault(k, v)
        ticketed, _ = sc_history()
        pending, seen = [], set()
        for u in [u for b in lb for u in b] + [u for b in rb for u in b]:
            if u not in seen and u not in ticketed:
                seen.add(u)
                pending.append(u)
        size = r.TICKET_BATCH_SIZE
        batches = [pending[i:i + size] for i in range(0, len(pending), size)]
        base = max([SWEEP_BATCH_BASE] + [b for b in sc_batch_numbers() if b > SWEEP_BATCH_BASE])
        log(f"  {len(pending):,} rule(s) not in any ticket -> {len(batches)} sweep batch(es)"
            + (f" (numbered {base + 1}-{base + len(batches)})" if batches else ""))
    else:
        batches = ledger_batches()
        base = 0
        log(f"Ledger has {len(batches)} recertify batch(es)")
        log("Loading mapping workbook and reporting-view owners...")
        mapping, owners = mapping_by_uid(), owners_by_rule()
        log(f"  {len(mapping):,} mapped rules, {len(owners):,} rules with an owner")
    todo = list(range(1, len(batches) + 1)) if BATCHES == "all" else list(BATCHES)
    for i in todo:
        if base and i > base:           # accept the displayed number too: 1077 == 77
            i -= base
        if not 1 <= i <= len(batches):
            log(f"Batch {i}: out of range (1-{len(batches)}) - skipped")
            continue
        n = base + i                    # the batch number recorded and shown (1001+ for reports)
        ticketed, open_tickets = sc_history()

        # 1. Finish any ticket from an earlier run of this batch before creating another.
        #    A run that resumes a ticket does not also create one for the same batch -
        #    run again for the rest, so each run does one clear thing per batch.
        mine = {tid: t for tid, t in open_tickets.items() if t["batch"] == n}
        #    In "create" mode open tickets are left for people to process, and the rules
        #    not yet in any ticket are created below - tracking is per rule, so no overlap.
        if mine and not DRY_RUN and STOP_AFTER == "create":
            log(f"\nBatch {n}: open ticket(s) {', '.join(mine)} left as they are - "
                f"creating the rest of the batch")
        elif mine and not DRY_RUN:
            for tid, t in mine.items():
                log(f"\nBatch {n}: resuming ticket {tid} (stage {t['stage']})")
                rules, _ = resolve_batch(t["uids"], mapping, owners)
                if not advance(tid, n, rules):
                    raise SystemExit("Stopped - fix the issue above before the next batch.")
            log(f"Batch {n}: resumed ticket(s) finished - run again to submit the rest of the batch")
            continue

        # 2. Only rules from this batch that are not already in one of our tickets.
        remaining = [u for u in batches[i - 1] if u not in ticketed]
        uids = remaining[:TEST_RULE_LIMIT or None]
        log(f"\nBatch {n}: {len(batches[i - 1])} rule(s) in batch, "
            f"{len(batches[i - 1]) - len(remaining)} already ticketed, submitting {len(uids)}")
        if not uids:
            continue

        rules, problems = resolve_batch(uids, mapping, owners)
        no_owner = sum(1 for x in rules if not x["owner"])
        log(f"  resolved {len(rules)}, problems {len(problems)}, without owner {no_owner}")
        for g, why in problems[:10]:
            log(f"    skip {g}: {why}")
        if not rules:
            log("  nothing to submit")
            continue

        payload = create_payload(n, rules)
        preview = OUT / f"sc_preview_batch{n}.json"
        preview.write_text(json.dumps({"create": payload,
                                       "decision": infos_block(rules),
                                       "rules": rules,
                                       "problems": problems}, indent=2), encoding="utf-8")
        log(f"  payload saved to {preview.name}")
        if DRY_RUN:
            continue

        resp = call("POST", f"{SC}/tickets/", data=json.dumps(payload))
        loc = resp.headers.get("Location", "")
        tid = loc.rstrip("/").split("/")[-1] if loc else None
        if not tid:
            raise SystemExit(f"Ticket created but no id returned (Location={loc!r}) - "
                             f"find it in SecureChange and add it to {SC_LEDGER.name} "
                             f"before re-running, or the rules will be submitted again.")
        sc_ledger({"batch": n, "ticket_id": tid, "stage": "created",
                   "count": len(rules), "uids": [x["graphql_id"] for x in rules]})
        log(f"  CREATED ticket {tid}")

        if STOP_AFTER == "create":
            log("  STOP_AFTER='create' - check the ticket in SecureChange")
            continue
        if not advance(tid, n, rules):
            raise SystemExit("Stopped - fix the issue above before the next batch.")

    if _UNRESOLVED:
        dbg = OUT / "unresolved_debug.json"
        dbg.write_text(json.dumps(_UNRESOLVED, indent=2), encoding="utf-8")
        log(f"\n{len(_UNRESOLVED)} rule(s) still unresolved - details in {dbg.name}")
    log("\nDone.")


if __name__ == "__main__":
    try:
        main()
    except RuntimeError as e:
        log(f"ERROR {e}")
        sys.exit(1)
